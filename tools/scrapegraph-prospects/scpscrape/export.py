"""Écriture des 4 fichiers xlsx au format Salesforce Inspector Reloaded."""

from __future__ import annotations

import csv
from datetime import date
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)

COLUMNS_VALIDES = [
    "Id", "Name", "Phone", "BillingPostalCode", "BillingCity", "Industry",
    "Clients__c", "Lead_Panier__c", "Lead_Confiance__c", "Lead_Signal__c",
    "Lead_Angle__c", "Lead_Decideur__c", "Lead_Decideur_Source__c",
    "LinkedIn_Entreprise_URL", "Lead_Decideur_LinkedIn_URL", "Description",
]

COLUMNS_A_CONFIRMER = [
    "Id", "Name", "Phone", "BillingPostalCode", "BillingCity", "Industry",
    "Lead_Panier__c", "Lead_Confiance__c", "Lead_Question_Verif__c",
    "Lead_Decideur__c", "Lead_Decideur_Source__c",
    "LinkedIn_Entreprise_URL", "Lead_Decideur_LinkedIn_URL", "Description",
]

COLUMNS_EXCLUS = [
    "Id", "Statut", "Name", "Phone", "BillingPostalCode", "BillingCity",
    "Industry", "Raison", "Description",
]

MAX_WIDTH = 60


def _write_sheet(path: Path, sheet_title: str, columns: List[str], rows: List[Dict[str, Any]]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title

    sheet.append(columns)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")

    for row in rows:
        sheet.append([_as_text(row.get(column, "")) for column in columns])

    for index, column in enumerate(columns, start=1):
        widest = max([len(column)] + [len(_as_text(r.get(column, ""))) for r in rows] or [len(column)])
        sheet.column_dimensions[get_column_letter(index)].width = min(widest + 2, MAX_WIDTH)

    sheet.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def _as_text(value: Any) -> str:
    """Tout est écrit en texte : les codes postaux gardent leur zéro initial."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value)


def _versioned(path: Path) -> Path:
    """Le mount OneDrive interdit l'écrasement : on suffixe _v2, _v3, ..."""
    if not path.exists():
        return path
    stem, suffix = path.stem, path.suffix
    version = 2
    while True:
        candidate = path.with_name(f"{stem}_v{version}{suffix}")
        if not candidate.exists():
            return candidate
        version += 1


def write_batch(outdir: Path, run_date: date, buckets: Dict[str, List[Dict[str, Any]]]) -> Dict[str, Path]:
    """Écrit les 4 xlsx du jour et retourne leurs chemins réels."""
    outdir = Path(outdir)
    stamp = run_date.isoformat()
    written: Dict[str, Path] = {}

    plan = [
        ("valides", f"{stamp}_VALIDES.xlsx", "Valides", COLUMNS_VALIDES),
        ("volume", f"{stamp}_VOLUME.xlsx", "Volume", COLUMNS_VALIDES),
        ("a_confirmer", f"{stamp}_A_CONFIRMER.xlsx", "A confirmer", COLUMNS_A_CONFIRMER),
        ("exclus", f"{stamp}_EXCLUS_ET_HORS_CIBLE.xlsx", "Exclus et Hors-cible", COLUMNS_EXCLUS),
    ]

    for key, filename, sheet_title, columns in plan:
        target = _versioned(outdir / filename)
        written[key] = _write_sheet(target, sheet_title, columns, buckets.get(key, []))

    return written


def write_stratif_csv(outdir: Path, run_date: date, scored) -> Path:
    """Trace interne : un CSV avec le score et le motif de chaque prospect."""
    path = _versioned(Path(outdir) / f"{run_date.isoformat()}_stratif_complete.csv")
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Name", "Website", "Panier", "Confiance", "Score", "Raison", "Signaux"])
        for record, verdict in scored:
            writer.writerow([
                record.get("Name", ""),
                record.get("Website", ""),
                verdict.panier,
                verdict.confiance,
                verdict.score,
                verdict.raison,
                verdict.signal_label,
            ])
    return path

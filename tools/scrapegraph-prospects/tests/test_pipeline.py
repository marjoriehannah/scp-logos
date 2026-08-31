"""Chaîne complète hors ligne : fixtures → 4 xlsx aux bonnes colonnes."""

import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest
from openpyxl import load_workbook

from scpscrape.cli import read_csv
from scpscrape.export import COLUMNS_EXCLUS, COLUMNS_VALIDES, write_batch, write_stratif_csv
from scpscrape.extract import FixtureExtractor, slugify, _coerce
from scpscrape.pipeline import format_decideur, parse_employees, process
from scpscrape.spec import load_spec

ROOT = Path(__file__).resolve().parents[1]
SPEC = load_spec(ROOT / "config" / "exemple-installateur.yaml")
RUN_DATE = date(2026, 8, 31)


@pytest.fixture(scope="module")
def result():
    rows = read_csv(ROOT / "fixtures" / "prospects.csv")
    extractor = FixtureExtractor(ROOT / "fixtures" / "extractions")
    return process(rows, extractor, SPEC, run_date=RUN_DATE)


def test_every_fixture_lands_in_its_expected_bucket(result):
    buckets = result["buckets"]
    paniers = {
        row["Name"]: row.get("Lead_Panier__c", row.get("Statut"))
        for bucket in buckets.values()
        for row in bucket
    }
    assert paniers["Hôtel Lac Panorama"] == "chaud"
    assert paniers["Clinique des Alpes"] == "tiede"
    assert paniers["Garage Mont-Blanc Automobiles"] == "tiede"
    assert paniers["Ibis Annecy Centre"] == "volume"
    assert paniers["Atelier Conseil Léman"] == "a_confirmer"
    assert paniers["Restaurant du Port"] == "a_confirmer"
    assert paniers["Miroiterie du Léman"] == "Exclu"
    assert paniers["Savoie Participations Holding"] == "Hors-cible"


def test_chaud_is_listed_first_in_valides(result):
    assert result["buckets"]["valides"][0]["Lead_Panier__c"] == "chaud"


def test_extraction_fills_missing_source_fields(result):
    hotel = next(r for r in result["buckets"]["valides"] if r["Name"] == "Hôtel Lac Panorama")
    assert hotel["Phone"] == "+33450264118"          # trouvé sur le site, formaté E.164
    assert hotel["BillingPostalCode"] == "74500"
    assert hotel["BillingCity"] == "Évian-les-Bains"
    assert hotel["Industry"] == "Hôtellerie"          # valeur de picklist, pas le mot du site
    assert hotel["Clients__c"] == SPEC.name
    assert hotel["Lead_Decideur__c"].startswith("Camille Reverdy")
    assert hotel["Lead_Decideur_Source__c"] == "ScrapeGraphAI (site web)"
    assert "printemps 2026" in hotel["Lead_Signal__c"]
    assert hotel["Description"].startswith("Batch Exemple SA 2026-08-31 — chaud")


def test_source_row_wins_over_extraction():
    """Une valeur fournie par Marjorie n'est jamais écrasée par le web."""
    rows = [{
        "Name": "Clinique des Alpes",
        "Website": "https://www.clinique-des-alpes.fr",
        "Phone": "+33499999999",
        "BillingPostalCode": "74000",
        "BillingCity": "Annecy",
        "BillingCountry": "France",
    }]
    out = process(rows, FixtureExtractor(ROOT / "fixtures" / "extractions"), SPEC, run_date=RUN_DATE)
    assert out["buckets"]["valides"][0]["Phone"] == "+33499999999"


def test_fallback_decideur_when_no_name_published(result):
    garage = next(r for r in result["buckets"]["valides"] if r["Name"].startswith("Garage"))
    assert garage["Lead_Decideur_Source__c"] == "Fallback (intitulé générique)"
    assert garage["Lead_Decideur__c"] == "Le chef des ventes / responsable du site"


def test_a_confirmer_carries_a_question(result):
    for row in result["buckets"]["a_confirmer"]:
        assert row["Lead_Question_Verif__c"]


def test_exclus_carry_statut_and_reason(result):
    for row in result["buckets"]["exclus"]:
        assert row["Statut"] in {"Exclu", "Hors-cible"}
        assert row["Raison"]
        assert len(row["Raison"]) <= 150


def test_dedup_excludes_known_account():
    rows = read_csv(ROOT / "fixtures" / "prospects.csv")
    out = process(
        rows,
        FixtureExtractor(ROOT / "fixtures" / "extractions"),
        SPEC,
        dedup_names={"hotel lac panorama"},
        run_date=RUN_DATE,
    )
    assert out["buckets"]["valides"] == [] or all(
        r["Name"] != "Hôtel Lac Panorama" for r in out["buckets"]["valides"]
    )
    assert any(r["Name"] == "Hôtel Lac Panorama" for r in out["buckets"]["exclus"])


def test_missing_fixture_does_not_break_the_batch():
    rows = [{"Name": "Site Inconnu", "Website": "https://www.jamais-vu.fr", "BillingPostalCode": "74000"}]
    out = process(rows, FixtureExtractor(ROOT / "fixtures" / "extractions"), SPEC, run_date=RUN_DATE)
    assert sum(len(b) for b in out["buckets"].values()) == 1


def test_extractor_failure_is_reported_not_fatal():
    class Broken:
        def extract(self, url, hint=None):
            raise RuntimeError("timeout playwright")

    rows = [{"Name": "Site Cassé", "Website": "https://www.casse.fr", "BillingPostalCode": "74000"}]
    out = process(rows, Broken(), SPEC, run_date=RUN_DATE)
    assert any("timeout playwright" in w for w in out["warnings"])
    assert sum(len(b) for b in out["buckets"].values()) == 1


def test_write_batch_produces_four_files_with_expected_headers(tmp_path, result):
    written = write_batch(tmp_path, RUN_DATE, result["buckets"])
    assert set(written) == {"valides", "volume", "a_confirmer", "exclus"}

    valides = load_workbook(written["valides"])
    assert valides.active.title == "Valides"
    assert [c.value for c in valides.active[1]] == COLUMNS_VALIDES
    assert valides.active.freeze_panes == "A2"

    exclus = load_workbook(written["exclus"])
    assert exclus.active.title == "Exclus et Hors-cible"
    assert [c.value for c in exclus.active[1]] == COLUMNS_EXCLUS


def test_postal_code_stays_text_in_xlsx(tmp_path):
    buckets = {"valides": [{"Name": "Ain Test", "BillingPostalCode": "01200"}], "volume": [], "a_confirmer": [], "exclus": []}
    written = write_batch(tmp_path, RUN_DATE, buckets)
    sheet = load_workbook(written["valides"]).active
    column = COLUMNS_VALIDES.index("BillingPostalCode") + 1
    assert sheet.cell(row=2, column=column).value == "01200"


def test_existing_file_is_versioned_not_overwritten(tmp_path, result):
    first = write_batch(tmp_path, RUN_DATE, result["buckets"])
    second = write_batch(tmp_path, RUN_DATE, result["buckets"])
    assert first["valides"].name == "2026-08-31_VALIDES.xlsx"
    assert second["valides"].name == "2026-08-31_VALIDES_v2.xlsx"
    assert first["valides"].exists()


def test_stratif_csv_traces_every_prospect(tmp_path, result):
    path = write_stratif_csv(tmp_path, RUN_DATE, result["scored"])
    lines = path.read_text(encoding="utf-8").splitlines()
    assert len(lines) == len(result["scored"]) + 1


def test_slugify_matches_fixture_names():
    assert slugify("https://www.hotel-lac-panorama.fr/") == "www-hotel-lac-panorama-fr"
    assert slugify("HTTP://Exemple.FR") == "exemple-fr"


def test_coerce_accepts_dict_json_and_garbage():
    assert _coerce({"company_name": "X"})["company_name"] == "X"
    assert _coerce('{"company_name": "X"}')["company_name"] == "X"
    assert _coerce('Voici le résultat : {"company_name": "X"} fin')["company_name"] == "X"
    assert _coerce("No answer found.") == {}
    assert _coerce(None) == {}


def test_parse_employees():
    assert parse_employees("28 collaborateurs") == 28
    assert parse_employees("équipe de 12 personnes") == 12
    assert parse_employees("une belle équipe") is None
    assert parse_employees(None) is None


def test_format_decideur_prefers_named_person():
    label, source = format_decideur({
        "Industry": "Hôtellerie",
        "_decision_makers": [{"full_name": "Camille Reverdy", "role": "Directrice", "email": "c@x.fr"}],
    })
    assert label == "Camille Reverdy — Directrice (c@x.fr)"
    assert source == "ScrapeGraphAI (site web)"
